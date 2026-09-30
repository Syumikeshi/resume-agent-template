#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
求职投递追踪表生成器 —— openpyxl 直写（无公式，避免依赖公式引擎）

用法：
    python3 build.py            # 增量更新（默认）：沿用旧表的过程数据
    python3 build.py --rebuild  # 全量重做：忽略旧表，从清单完全重建

输入：tracker_data.py（由 extract_tracker_data.py 从 lists/ 清单自动生成，勿手改）
输出：追踪表.xlsx

生成的工作表：
  1. 招聘会现场记录 —— 空白模板（带下拉校验）；增量模式下**原样保留**已填内容
  2. 每份清单一个方向表，统一十列：
     序号 | 优先级 | 企业 / 机构 | WLB | 岗位 / 方向 | 地点 |
     薪资（含核实状态）| 投递日期 | 进展 | 备注

增量更新规则（默认）：
  - 按「工作表名 + 企业名」匹配，沿用旧表的 **投递日期** 与 **进展** 两列
  - 清单里新增的企业按「待投递」进入；结构性内容（梯队/薪资/备注）以清单为准，会被覆盖
  - 旧表中有、新清单中没有的方向表会被移除（清单是唯一真相源，删清单=删方向）
  - 招聘会现场记录整表保留，行数不够自动扩展
"""

import os
import sys

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

try:
    from tracker_data import DATA
except ImportError:
    raise SystemExit(
        "缺少 tracker_data.py —— 请先运行：python3 extract_tracker_data.py")

OUTPUT = "追踪表.xlsx"
WORKBOOK_TITLE = "求职投递追踪表"
REBUILD = "--rebuild" in sys.argv
FAIR_SHEET = "招聘会现场记录"


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEADER_BG = xl_color("#4472C4")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_BORDER = xl_color("#BFBFBF")
XL_LIGHT = xl_color("#F2F7FF")
XL_GREEN_BG = xl_color("#C6EFCE"); XL_GREEN_FG = xl_color("#006100")
XL_RED_BG   = xl_color("#FFC7CE"); XL_RED_FG   = xl_color("#9C0006")
XL_AMBER_BG = xl_color("#FFEB9C"); XL_AMBER_FG = xl_color("#9C6500")
XL_TIER1_BG = xl_color("#E2F0D9"); XL_TIER1_FG = xl_color("#1F5C1F")
XL_TIER2_BG = xl_color("#FFF2CC"); XL_TIER2_FG = xl_color("#7F6000")

F_HEADER = Font(name="Microsoft YaHei", size=10, bold=True, color=XL_HEADER_FG)
F_BODY = Font(name="Microsoft YaHei", size=10)
AL_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
BORDER = Border(
    left=Side(style="thin", color=XL_BORDER),
    right=Side(style="thin", color=XL_BORDER),
    top=Side(style="thin", color=XL_BORDER),
    bottom=Side(style="thin", color=XL_BORDER),
)

STATUS = '"待投递,已投递,已笔试,面试中,已录用,已拒绝,无回复"'
PLAN_HEADERS = ["序号", "优先级", "企业 / 机构", "WLB", "岗位 / 方向", "地点",
                "薪资（含核实状态）", "投递日期", "进展", "备注"]
PLAN_WIDTHS = [6, 12, 26, 8, 36, 18, 32, 12, 12, 60]
BLANK_ROWS = 30   # 现场记录表预留的可填写行数（增量模式下按需扩展）


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = F_HEADER
        cell.fill = PatternFill("solid", fgColor=XL_HEADER_BG)
        cell.alignment = AL_CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def style_body(ws, r1, r2, ncols):
    for r in range(r1, r2 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = F_BODY
            cell.border = BORDER
            cell.alignment = AL_LEFT
            cell.fill = PatternFill("solid", fgColor=XL_LIGHT)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def load_existing(path):
    """读旧表。返回 (各方向表的过程数据 {表名: {企业: (投递日期, 进展)}}, 现场记录已填行)."""
    if REBUILD or not os.path.exists(path):
        return {}, []
    wb_old = openpyxl.load_workbook(path, data_only=True)
    process, fair_rows = {}, []
    if FAIR_SHEET in wb_old.sheetnames:
        for row in wb_old[FAIR_SHEET].iter_rows(min_row=2, values_only=True):
            # 序号列(A)默认预填，只有 A 列以外有内容才算用户填过
            if any(v is not None and str(v).strip() != "" for v in row[1:]):
                fair_rows.append(list(row))
    for ws in wb_old.worksheets:
        if ws.title == FAIR_SHEET:
            continue
        d = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) > 8 and row[2]:
                d[str(row[2]).strip()] = (row[7], row[8])
        process[ws.title] = d
    wb_old.close()
    return process, fair_rows


wb = Workbook()

# ============================================================
# Sheet 1 —— 招聘会现场记录（空白模板；增量模式回填旧内容）
# ============================================================
process, fair_rows = load_existing(OUTPUT)

ws1 = wb.active
ws1.title = FAIR_SHEET

H1 = ["序号", "日期", "企业", "岗位 / 方向", "前台还是中后台",
      "有无业绩指标", "HR 最看重的能力（原话）", "HR 联系方式",
      "已递交简历版本", "下一步动作", "备注"]
W1 = [6, 12, 22, 26, 15, 13, 30, 20, 20, 22, 24]
LAST = 1 + max(BLANK_ROWS, len(fair_rows))

for i, h in enumerate(H1, start=1):
    ws1.cell(row=1, column=i, value=h)
style_header(ws1, 1, len(H1))
for i in range(2, LAST + 1):
    ws1.cell(row=i, column=1, value=i - 1)
style_body(ws1, 2, LAST, len(H1))
for r, old_row in enumerate(fair_rows, start=2):
    for c, v in enumerate(old_row, start=1):
        if v is not None:
            ws1.cell(row=r, column=c, value=v)
for i in range(2, LAST + 1):
    ws1.cell(row=i, column=1).alignment = AL_CENTER
    ws1.cell(row=i, column=2).alignment = AL_CENTER
    ws1.cell(row=i, column=2).number_format = "YYYY-MM-DD"
    ws1.cell(row=i, column=5).alignment = AL_CENTER
    ws1.cell(row=i, column=6).alignment = AL_CENTER
set_widths(ws1, W1)
ws1.freeze_panes = "C2"
ws1.auto_filter.ref = f"A1:K{LAST}"

dv_scope1 = DataValidation(type="list", formula1='"前台,中后台,不确定"', allow_blank=True)
ws1.add_data_validation(dv_scope1); dv_scope1.add(f"E2:E{LAST}")
dv_metric1 = DataValidation(type="list", formula1='"有,无,不确定"', allow_blank=True)
ws1.add_data_validation(dv_metric1); dv_metric1.add(f"F2:F{LAST}")


# ============================================================
# 方向表构建器：接收 tracker_data.DATA 的条目，统一写出十列
# ============================================================
def build_plan_sheet(sheet_name, rows, carry):
    """carry: {企业名: (投递日期, 进展)}——增量模式下沿用旧表过程数据。"""
    ws = wb.create_sheet(sheet_name)
    ws.append(PLAN_HEADERS)
    style_header(ws, 1, len(PLAN_HEADERS))
    kept = 0
    for idx, r in enumerate(rows, 1):
        old = carry.get(str(r["company"]).strip())
        date_v = old[0] if old and old[0] is not None else ""
        status_v = old[1] if old and old[1] else "待投递"
        if old and (date_v != "" or (old[1] and old[1] != "待投递")):
            kept += 1
        ws.append([idx, r["tier"], r["company"], r["wlb"], r["role"],
                   r["location"], r["salary"], date_v, status_v, r["note"]])
    last = 1 + len(rows)
    style_body(ws, 2, last, len(PLAN_HEADERS))
    for i in range(2, last + 1):
        ws.cell(row=i, column=1).alignment = AL_CENTER
        ws.cell(row=i, column=2).alignment = AL_CENTER
        ws.cell(row=i, column=4).alignment = AL_CENTER
        ws.cell(row=i, column=8).number_format = "YYYY-MM-DD"
        ws.cell(row=i, column=9).alignment = AL_CENTER
    set_widths(ws, PLAN_WIDTHS)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:J{last}"

    # 进展列：下拉 + 状态着色
    dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"I2:I{last}")
    for val, bg, fg in [("已录用", XL_GREEN_BG, XL_GREEN_FG),
                        ("面试中", XL_AMBER_BG, XL_AMBER_FG),
                        ("已笔试", XL_AMBER_BG, XL_AMBER_FG),
                        ("已拒绝", XL_RED_BG, XL_RED_FG),
                        ("无回复", XL_RED_BG, XL_RED_FG)]:
        ws.conditional_formatting.add(
            f"I2:I{last}",
            CellIsRule(operator="equal", formula=[f'"{val}"'],
                       fill=PatternFill("solid", fgColor=bg), font=Font(color=fg, size=10)))
    # 优先级列：按梯队着色
    for tier, bg, fg in (("第一梯队", XL_TIER1_BG, XL_TIER1_FG),
                         ("第二梯队", XL_TIER2_BG, XL_TIER2_FG)):
        ws.conditional_formatting.add(
            f"B2:B{last}",
            CellIsRule(operator="equal", formula=[f'"{tier}"'],
                       fill=PatternFill("solid", fgColor=bg),
                       font=Font(color=fg, size=10, bold=(tier == "第一梯队"))))
    return ws, kept


# ============================================================
# 每份清单一个方向表（增量合并过程数据）
# ============================================================
report = []
for sheet_name, rows in DATA.items():
    _, kept = build_plan_sheet(sheet_name, rows, process.get(sheet_name, {}))
    report.append((sheet_name, len(rows), kept))

# 旧表里有、新清单里没有的方向 → 警告（数据随表移除）
for old_name in process:
    if old_name not in DATA:
        print(f"⚠️  旧表的「{old_name}」在 lists/ 清单里已不存在，该表（含过程数据）已移除")

wb.properties.title = WORKBOOK_TITLE
wb.save(OUTPUT)

_total = sum(n for _, n, _ in report)
mode = "全量重做（--rebuild）" if REBUILD else "增量更新"
print(f"已生成 {OUTPUT}（{mode}）：1 个现场记录表 + {len(DATA)} 个方向表，共 {_total} 条目标")
for name, n, kept in report:
    extra = f"，沿用旧表过程数据 {kept} 条" if (kept and not REBUILD) else ""
    print(f"  {name:<14} {n:>3} 条{extra}")
if fair_rows and not REBUILD:
    print(f"  现场记录表已保留 {len(fair_rows)} 条手填记录")
print("  要彻底重做（忽略旧表过程数据）：python3 build.py --rebuild")
